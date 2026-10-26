
import io
import json
from pathlib import Path
from uuid import uuid4

import pandas as pd
from docx import Document
from fastapi import UploadFile
from pypdf import PdfReader


ALLOWED_EXTENSIONS = {
    ".pdf",
    ".xlsx",
    ".xls",
    ".docx",
    ".txt",
    ".csv",
}


async def ingest_document(file: UploadFile) -> dict:
    """
    Read an uploaded document, extract its text,
    and return JSON-compatible data for processing.
    """

    if not file.filename:
        raise ValueError("No file was provided.")

    filename = Path(file.filename).name
    extension = Path(filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. Upload PDF, Excel, CSV, Word, or TXT."
        )

    document_id = str(uuid4())

    try:
        file_bytes = await file.read()

        if not file_bytes:
            raise ValueError("The uploaded file is empty.")

        # 1. PDF FILE
        if extension == ".pdf":
            reader = PdfReader(io.BytesIO(file_bytes))

            if reader.is_encrypted:
                raise ValueError(
                    "Encrypted PDFs are not supported."
                )

            pages = []
            text_parts = []

            for page_number, page in enumerate(
                reader.pages, start=1
            ):
                page_text = page.extract_text() or ""

                pages.append({
                    "page": page_number,
                    "text": page_text,
                })

                text_parts.append(page_text)

            extracted_text = "\n".join(text_parts)

            metadata = {
                "page_count": len(pages),
                "pages": pages,
            }

        # 2. EXCEL FILES
        elif extension in {".xlsx", ".xls"}:
            sheets = pd.read_excel(
                io.BytesIO(file_bytes),
                sheet_name=None,
            )

            sheet_results = []
            text_parts = []

            for sheet_name, dataframe in sheets.items():
                dataframe = dataframe.fillna("")

                sheet_text = dataframe.to_csv(
                    index=False,
                    sep="\t",
                )

                sheet_results.append({
                    "sheet_name": str(sheet_name),
                    "text": sheet_text,
                })

                text_parts.append(
                    f"Sheet: {sheet_name}\n{sheet_text}"
                )

            extracted_text = "\n\n".join(text_parts)

            metadata = {
                "sheet_count": len(sheet_results),
                "sheets": sheet_results,
            }

        # 3. CSV FILES
        elif extension == ".csv":
            try:
                dataframe = pd.read_csv(
                    io.BytesIO(file_bytes),
                    encoding="utf-8-sig",
                )
            except UnicodeDecodeError:
                dataframe = pd.read_csv(
                    io.BytesIO(file_bytes),
                    encoding="cp1252",
                )

            # Replace missing values with empty strings.
            dataframe = dataframe.fillna("")

            # Convert CSV rows into text for later chunking.
            extracted_text = dataframe.to_csv(
                index=False,
                sep="\t",
            )

            metadata = {
                "row_count": len(dataframe),
                "column_count": len(dataframe.columns),
                "columns": [
                    str(column)
                    for column in dataframe.columns
                ],
            }

        # 4. WORD DOCUMENTS
        elif extension == ".docx":
            document = Document(io.BytesIO(file_bytes))

            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]

            table_text = []

            for table in document.tables:
                for row in table.rows:
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    table_text.append(" | ".join(cells))

            extracted_text = "\n".join(paragraphs)

            if table_text:
                extracted_text += (
                    "\n\n" + "\n".join(table_text)
                )

            metadata = {
                "paragraph_count": len(paragraphs),
                "table_count": len(document.tables),
            }

        # 5. TEXT FILES
        elif extension == ".txt":
            try:
                extracted_text = file_bytes.decode("utf-8-sig")
            except UnicodeDecodeError as error:
                raise ValueError(
                    "The TXT file must use UTF-8 text encoding."
                ) from error

            metadata = {}

        else:
            raise ValueError(
                "Unsupported file type."
            )

        # Check that the document contains readable text.
        if not extracted_text.strip():
            raise ValueError(
                "No readable text was found in the uploaded document."
            )

        # Build the ingestion result for the next pipeline stage.
        result = {
            "document_id": document_id,
            "filename": filename,
            "file_type": extension,
            "status": "success",
            "text": extracted_text,
            "metadata": metadata,
        }

        # Verify that the result is JSON-compatible.
        json.dumps(result, ensure_ascii=False)

        return result

    except ValueError:
        raise

    except Exception as error:
        raise ValueError(
            f"Could not process the uploaded file: {error}"
        ) from error

    finally:
        await file.close()
